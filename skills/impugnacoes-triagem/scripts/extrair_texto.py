#!/usr/bin/env python3
"""Extrai o texto dos documentos do edital com marcadores de página, parágrafo ou célula. Python 3.10+.

Uso:
  extrair_texto.py ARQUIVO... --saida DIR [--max-mb 200]

Para cada arquivo grava DIR/<nome-ascii>.<ext>.txt e atualiza DIR/inventario.json. Formatos: PDF (pypdf),
DOCX, XLSX (openpyxl), TXT e MD. Imprime em JSON o resumo desta execução e sai com 1 se algum arquivo
falhou; os demais são gravados mesmo assim.
"""
import argparse
import codecs
import hashlib
import json
import logging
import math
import os
import re
import sys
import tempfile
import unicodedata
import warnings
import zipfile
from datetime import datetime, time
from pathlib import Path
from xml.etree import ElementTree

MB = 1024 * 1024
ZIP_RATIO = 5
INVENTORY = 'inventario.json'
KINDS = {'.pdf': 'pdf', '.docx': 'docx', '.xlsx': 'xlsx', '.txt': 'texto', '.md': 'texto'}
INSTALL = 'instale com .venv/bin/pip install -r scripts/requirements.txt'
NATIVE = 'leia o PDF com a leitura nativa do Claude Code (Read com pages)'
NO_PYPDF = f'pypdf não está instalado: {INSTALL} ou {NATIVE}'
NO_CRYPTO = f'PDF cifrado com AES exige o pacote cryptography, que não está instalado: {INSTALL} ou {NATIVE}'
NO_TEXT = '[página sem texto extraível: pode exigir OCR ou leitura visual]'
PAGE_FAILED = '[falha ao extrair esta página'
# Linha do documento que imitaria um marcador do .txt; ganha um espaço na frente.
FAKE_MARKER = re.compile(r'(?m)^(?====|\[(?:¶|página sem texto|falha ao extrair))')
DOCX_HEADER = ['# Paginação: não verificada (DOCX)',
               '# Parágrafos: [¶N] numera só os não vazios, na ordem do corpo, incluindo células de tabela; '
               'a numeração automática de listas do Word não aparece no texto']
XLSX_HEADER = ['# Valores: em cache da planilha; fórmulas não recalculadas']
BAD_ZIP = 'arquivo corrompido ou não é um documento Office (zip) válido'
BAD_DATE = re.compile(r'Cell (\w+) is marked as a date')
BOMS = [(codecs.BOM_UTF32_LE, 'utf-32'), (codecs.BOM_UTF32_BE, 'utf-32'),
        (codecs.BOM_UTF16_LE, 'utf-16'), (codecs.BOM_UTF16_BE, 'utf-16')]  # UTF-32 LE começa como UTF-16 LE
SURROGATE = re.compile('[\ud800-\udfff]')

W_NS = ('http://schemas.openxmlformats.org/wordprocessingml/2006/main',
        'http://purl.oclc.org/ooxml/wordprocessingml/main')
MC_FALLBACK = '{http://schemas.openxmlformats.org/markup-compatibility/2006}Fallback'
RUN_PIECES = {'tab': ' ', 'br': ' ', 'cr': ' ', 'noBreakHyphen': '-'}


class ExtractionError(Exception):
    pass


def nfc(text):
    return unicodedata.normalize('NFC', text)


def printable(text):
    """Texto gravável em UTF-8: bytes de nome fora do UTF-8, que o Python guarda como surrogates, viram U+FFFD."""
    return SURROGATE.sub('\ufffd', text)


def file_sha256(path):
    digest = hashlib.sha256()
    with open(path, 'rb') as fh:
        for block in iter(lambda: fh.read(MB), b''):
            digest.update(block)
    return digest.hexdigest()


def clean(text):
    return nfc(text.replace('\r\n', '\n').replace('\r', '\n').replace('\x00', '')).strip()


def one_line(text):
    return ' '.join(nfc(text).split())


def slug(text):
    text = unicodedata.normalize('NFKD', text).encode('ascii', 'ignore').decode('ascii').lower()
    return re.sub(r'[^a-z0-9]+', '-', text).strip('-')


def output_name(name, n=1):
    """Nome ASCII do .txt: "Edital PE 1.2026.pdf" vira "edital-pe-1-2026.pdf.txt"; na colisão, "-2" antes da extensão."""
    path = Path(nfc(name))
    base = slug(path.stem)[:150].strip('-') or 'documento'
    if n > 1:
        base = f'{base}-{n}'
    ext = slug(path.suffix)
    return f'{base}.{ext}.txt' if ext else f'{base}.txt'


def locate(path):
    """O caminho como veio; se não existir, procura cada parte pelo nome em NFC (pastas podem estar em NFD).

    Assim o caminho do inventário, gravado em NFC, também abre o arquivo."""
    if path.exists():
        return checked_file(path)
    current, parts = (Path(path.anchor), path.parts[1:]) if path.is_absolute() else (Path('.'), path.parts)
    for part in parts:
        candidate = current / part
        if not candidate.exists():
            matches = sorted(e for e in current.iterdir() if nfc(e.name) == nfc(part)) if current.is_dir() else []
            if not matches:
                raise ExtractionError('arquivo não encontrado')
            candidate = matches[0]
        current = candidate
    return checked_file(current)


def checked_file(path):
    if not path.is_file():
        raise ExtractionError('não é um arquivo')
    return path


def check_limit(max_mb):
    if not math.isfinite(max_mb) or max_mb <= 0:
        raise ValueError('--max-mb deve ser um número positivo e finito')


def check_size(path, max_mb):
    size = path.stat().st_size
    if not size:
        raise ExtractionError('arquivo vazio (0 byte): baixe ou copie o arquivo de novo')
    if size > max_mb * MB:
        raise ExtractionError(f'arquivo de {size / MB:.1f} MB é maior que o limite de {max_mb:g} MB (--max-mb)')


def check_zip(path, max_mb):
    """Recusa DOCX/XLSX cujo conteúdo descompactado passe de ZIP_RATIO vezes o limite (zip bomb)."""
    try:
        with zipfile.ZipFile(path) as zf:
            total = sum(info.file_size for info in zf.infolist())
    except zipfile.BadZipFile:
        raise ExtractionError(BAD_ZIP) from None
    if total > ZIP_RATIO * max_mb * MB:
        raise ExtractionError(f'conteúdo descompactado de {total / MB:.1f} MB passa do limite de '
                              f'{ZIP_RATIO * max_mb:g} MB (possível zip bomb)')


def pdf_pages(path):
    """Texto limpo de cada página: vazio se não houver texto, marcador de falha se a extração quebrar."""
    try:
        from pypdf import PasswordType, PdfReader
        from pypdf.errors import DependencyError, PdfReadError
    except ImportError:
        raise ExtractionError(NO_PYPDF) from None
    logging.getLogger('pypdf').setLevel(logging.ERROR)
    source = locate(Path(path))
    try:
        reader = PdfReader(source)  # já tenta a senha vazia, e no AES-256 isso exige o cryptography
        if reader.is_encrypted and reader.decrypt('') == PasswordType.NOT_DECRYPTED:
            raise ExtractionError('PDF protegido por senha de abertura')
        return [page_text(page) for page in reader.pages]
    except DependencyError:
        raise ExtractionError(NO_CRYPTO) from None
    except PdfReadError as exc:
        raise ExtractionError(f'não é um PDF válido ou está corrompido ({type(exc).__name__}): baixe o arquivo '
                              f'de novo ou {NATIVE}') from None


def page_text(page):
    from pypdf.errors import DependencyError
    try:
        return FAKE_MARKER.sub(' ', clean(page.extract_text() or ''))
    except DependencyError:
        raise  # falta de pacote vale para o arquivo inteiro: a página não está sem texto
    except Exception as exc:  # uma página malformada não derruba o documento inteiro
        return f'{PAGE_FAILED} ({type(exc).__name__}): pode exigir leitura visual]'


def blank_pages(pages):
    return [n for n, text in enumerate(pages, 1) if not text or text.startswith(PAGE_FAILED)]


def render_pdf(pages):
    blocks = [f'=== página PDF {n} ===\n{text or NO_TEXT}' for n, text in enumerate(pages, 1)]
    return '\n\n'.join(blocks) + '\n'


def extract_pdf(path):
    """(texto com marcadores, número de páginas, páginas sem texto)."""
    pages = pdf_pages(path)
    return render_pdf(pages), len(pages), blank_pages(pages)


def w_tag(element):
    """Nome local de um elemento do WordprocessingML; None para outros namespaces."""
    namespace, _, name = element.tag.rpartition('}')
    return name if namespace.lstrip('{') in W_NS else None


def walk(element, skip=()):
    """Descendentes em ordem de documento, sem entrar nos ramos cujo nome local está em skip."""
    for child in element:
        if child.tag == MC_FALLBACK:  # cópia alternativa do mesmo conteúdo (caixas de texto)
            continue
        yield child
        if w_tag(child) not in skip:
            yield from walk(child, skip)


def paragraph_text(paragraph):
    # O texto das caixas de texto aninhadas sai como parágrafo próprio, logo depois deste.
    parts = []
    for node in walk(paragraph, skip=('txbxContent',)):
        name = w_tag(node)
        if name == 't':
            parts.append(node.text or '')
        elif name in RUN_PIECES:
            parts.append(RUN_PIECES[name])
    return one_line(''.join(parts))


def docx_part(zf):
    """Parte principal do documento, pelo relacionamento officeDocument (quase sempre word/document.xml)."""
    try:
        rels = ElementTree.fromstring(zf.read('_rels/.rels'))
    except (KeyError, ElementTree.ParseError):
        return 'word/document.xml'
    for rel in rels:
        if rel.get('Type', '').endswith('/officeDocument'):
            return rel.get('Target', '').lstrip('/')
    return 'word/document.xml'


def docx_paragraphs(path, max_mb=200):
    path = locate(Path(path))
    check_zip(path, max_mb)
    with zipfile.ZipFile(path) as zf:
        try:
            fh = zf.open(docx_part(zf))
        except KeyError:
            raise ExtractionError('DOCX sem o corpo do documento (word/document.xml)') from None
        # O zipfile para no tamanho declarado, que o check_zip já limitou; se ele mentir, o CRC falha.
        try:
            with fh:
                return stream_paragraphs(fh)
        except zipfile.BadZipFile:
            raise ExtractionError(BAD_ZIP) from None


def stream_paragraphs(source):
    """Textos não vazios dos w:p em ordem de documento, lendo o XML em fluxo.

    Só o parágrafo aberto fica na memória. O de uma caixa de texto fecha antes do parágrafo que o contém,
    por isso cada um é inserido na posição em que abriu."""
    texts, starts, stack, fallback = [], [], [], 0
    for event, element in ElementTree.iterparse(source, events=('start', 'end')):
        is_paragraph = w_tag(element) == 'p' and not fallback
        if event == 'start':
            fallback += element.tag == MC_FALLBACK
            if is_paragraph:
                starts.append(len(texts))
            stack.append(element)
            continue
        stack.pop()
        fallback -= element.tag == MC_FALLBACK
        if is_paragraph:
            start = starts.pop()
            if text := paragraph_text(element):
                texts.insert(start, text)
        if stack and not starts:
            stack[-1].remove(element)
    return texts


def render_docx(paragraphs):
    return ''.join(f'[¶{n}] {text}\n' for n, text in enumerate(paragraphs, 1))


def extract_docx(path, max_mb=200):
    """Parágrafos não vazios do corpo, cada um como "[¶N] texto"."""
    return render_docx(docx_paragraphs(path, max_mb))


def cell_value(value):
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == time() else value.isoformat(sep=' ')
    if hasattr(value, 'isoformat'):
        return value.isoformat()
    return one_line(str(value))


def xlsx_sheets(path, max_mb=200):
    """[(nome, oculta, [(linha, [(coluna, valor)])], [células com data inválida])], só com as não vazias."""
    path = locate(Path(path))
    check_zip(path, max_mb)
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise ExtractionError(f'openpyxl não está instalado: {INSTALL}') from None
    with warnings.catch_warnings(record=True):  # avisos de estilo e extensões não afetam os valores
        warnings.simplefilter('always')
        workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        return [(ws.title, ws.sheet_state != 'visible', *sheet_rows(ws)) for ws in workbook.worksheets]
    finally:
        workbook.close()


def sheet_rows(ws):
    ws.reset_dimensions()  # alguns geradores gravam dimensões erradas, e o modo read_only cortaria linhas
    rows = []
    with warnings.catch_warnings(record=True) as caught:  # o read_only lê as células só aqui
        warnings.simplefilter('always')
        for row in ws.iter_rows():
            cells = [(c.row, c.column_letter, cell_value(c.value)) for c in row if c.value is not None]
            cells = [cell for cell in cells if cell[2]]
            if cells:
                rows.append((cells[0][0], [(col, value) for _, col, value in cells]))
    # O openpyxl troca por #VALUE! a data fora do intervalo do Excel e só avisa.
    bad_dates = [m.group(1) for w in caught if (m := BAD_DATE.match(str(w.message)))]
    return rows, bad_dates


def bad_dates_note(cells):
    shown = ', '.join(cells[:10]) + (f' e mais {len(cells) - 10}' if len(cells) > 10 else '')
    return f'[aviso: data fora do intervalo do Excel em {shown}; aparece como #VALUE!, que não vem da planilha]'


def render_xlsx(sheets):
    blocks = []
    for name, hidden, rows, bad_dates in sheets:
        lines = [f'=== aba: {name} ===']
        if hidden:
            lines.append('[aba oculta]')
        if bad_dates:
            lines.append(bad_dates_note(bad_dates))
        lines += [f'L{n}: ' + ' | '.join(f'{col}={value}' for col, value in cells) for n, cells in rows]
        if not rows:
            lines.append('[aba sem células preenchidas]')
        blocks.append('\n'.join(lines))
    return '\n\n'.join(blocks) + '\n'


def extract_xlsx(path, max_mb=200):
    """Por aba, cada linha com conteúdo como "L<n>: A=valor | C=valor"."""
    return render_xlsx(xlsx_sheets(path, max_mb))


def read_text(path):
    data = Path(path).read_bytes()
    for bom, encoding in BOMS:
        if data.startswith(bom):
            return clean(data.decode(encoding, errors='replace')), encoding
    if data.count(b'\x00') > len(data) // 4:
        raise ExtractionError('texto com muitos bytes nulos, provavelmente UTF-16 sem BOM: salve o arquivo '
                              'como UTF-8 e rode de novo')
    try:
        return clean(data.decode('utf-8-sig')), 'utf-8'
    except UnicodeDecodeError:
        return clean(data.decode('cp1252', errors='replace')), 'cp1252'


def extract(kind, path, max_mb):
    """(corpo do .txt, linhas extras do cabeçalho, metadados do inventário). Caracteres não contam marcadores."""
    if kind == 'pdf':
        pages = pdf_pages(path)
        blank = blank_pages(pages)
        chars = sum(len(text) for n, text in enumerate(pages, 1) if n not in blank)
        header = [f'# Páginas: {len(pages)}', f"# Páginas sem texto: {', '.join(map(str, blank)) or 'nenhuma'}"]
        return render_pdf(pages), header, {'paginas': len(pages), 'paginas_sem_texto': blank, 'caracteres': chars}
    if kind == 'docx':
        paragraphs = docx_paragraphs(path, max_mb)
        return render_docx(paragraphs), DOCX_HEADER, {'caracteres': sum(map(len, paragraphs))}
    if kind == 'xlsx':
        sheets = xlsx_sheets(path, max_mb)
        chars = sum(len(value) for _, _, rows, _ in sheets for _, cells in rows for _, value in cells)
        return render_xlsx(sheets), XLSX_HEADER, {'caracteres': chars}
    text, encoding = read_text(path)
    return f'=== texto ===\n{text}\n', [f'# Codificação: {encoding}'], {'caracteres': len(text)}


def write_atomic(path, data):
    fd, tmp = tempfile.mkstemp(dir=path.parent, prefix=f'.{path.name}.', suffix='.tmp')
    try:
        with os.fdopen(fd, 'wb') as fh:
            fh.write(data)
        os.replace(tmp, path)
    except BaseException:
        Path(tmp).unlink(missing_ok=True)
        raise


def load_inventory(out_dir):
    path = out_dir / INVENTORY
    if not path.exists():
        return {}
    try:
        items = json.loads(path.read_text(encoding='utf-8'))
        inventory = {nfc(item['arquivo']): item for item in items}
        bad = next((item['saida'] for item in items if not valid_output(item.get('saida'))), None)
    except (ValueError, TypeError, KeyError, AttributeError) as exc:
        raise ValueError(f'{path} inválido ({exc}); corrija ou apague o arquivo para recriá-lo') from None
    if bad is not None:
        raise ValueError(f"{path} inválido ('saida' {bad!r} não é um nome .txt simples dentro da pasta); "
                         'corrija ou apague o arquivo para recriá-lo')
    return inventory


def valid_output(name):
    """'saida' do inventário: None ou um .txt direto na pasta, para nunca gravar nem apagar fora dela."""
    return name is None or (isinstance(name, str) and Path(name).name == name and name.endswith('.txt'))


def save_inventory(out_dir, inventory):
    items = [inventory[key] for key in sorted(inventory)]
    write_atomic(out_dir / INVENTORY, (json.dumps(items, ensure_ascii=False, indent=2) + '\n').encode('utf-8'))


def assign_name(name, key, inventory, taken):
    """Reaproveita a saída anterior do mesmo arquivo; senão, o primeiro nome livre (-2, -3...)."""
    previous = inventory.get(key, {}).get('saida')
    others = {item['saida'] for k, item in inventory.items() if k != key and item.get('saida')}
    if previous and previous not in taken and previous not in others:
        return previous
    n = 1
    while output_name(name, n) in taken | others:
        n += 1
    return output_name(name, n)


def new_entry(key, name):
    return {'arquivo': key, 'nome': name, 'sha256': None, 'tipo': KINDS.get(Path(name).suffix.lower()),
            'paginas': None, 'paginas_sem_texto': None, 'caracteres': None, 'saida': None, 'erro': None}


def process_one(raw, key, out_dir, max_mb, inventory, taken):
    entry = new_entry(key, nfc(printable(Path(raw).name)))
    try:
        if SURROGATE.search(os.path.abspath(raw)):
            raise ExtractionError('o caminho tem bytes fora do UTF-8 (nome gravado em outra codificação): '
                                  'renomeie o arquivo ou a pasta e rode de novo')
        source = locate(Path(raw))
        check_size(source, max_mb)
        entry['sha256'] = file_sha256(source)
        if entry['tipo'] is None:
            raise ExtractionError(f"formato não suportado ({Path(raw).suffix or 'sem extensão'}): "
                                  'use PDF, DOCX, XLSX, TXT ou MD')
        body, extra, info = extract(entry['tipo'], source, max_mb)
        header = [f"# Fonte: {entry['nome']}", f"# SHA-256: {entry['sha256']}", f"# Tipo: {entry['tipo']}", *extra]
        entry.update(info, saida=assign_name(entry['nome'], key, inventory, taken))
        write_atomic(out_dir / entry['saida'], ('\n'.join(header) + '\n\n' + body).encode('utf-8'))
        taken.add(entry['saida'])
    except ExtractionError as exc:
        entry.update(erro=str(exc), saida=None)
    except ElementTree.ParseError as exc:
        entry.update(erro=f'XML interno do documento inválido, corrompido ou malicioso ({exc}): abra e salve '
                          'o arquivo de novo no editor de origem', saida=None)
    except Exception as exc:  # PDF ou planilha malformados lançam erros variados da biblioteca
        entry.update(erro=f'falha ao ler o arquivo ({type(exc).__name__}): {printable(str(exc))}', saida=None)
    remove_stale(out_dir, inventory.get(key, {}).get('saida'), taken)
    return entry


def remove_stale(out_dir, old, taken):
    """Apaga o texto de uma extração anterior que não vale mais (por exemplo, o arquivo agora falha)."""
    if old and old not in taken and Path(old).name == old:
        (out_dir / old).unlink(missing_ok=True)


def same_file(first, second):
    try:
        return os.path.samefile(locate(Path(first)), locate(Path(second)))
    except (OSError, ExtractionError):
        return True  # nenhum dos dois abre, e o primeiro já registrou o erro


def twin_entry(key, raw):
    """Arquivo diferente com o mesmo caminho em NFC: fica só no resumo, para não apagar o primeiro do inventário."""
    entry = new_entry(key, nfc(printable(Path(raw).name)))
    entry['erro'] = ('outro arquivo com o mesmo nome em NFC (só a forma dos acentos muda) já foi processado nesta '
                     'execução: renomeie um dos dois e rode de novo')
    return entry


def process(paths, out_dir, max_mb=200):
    """Extrai cada arquivo, grava os .txt e o inventário e devolve as entradas desta execução."""
    check_limit(max_mb)
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    inventory = load_inventory(out_dir)
    entries, taken, seen = [], set(), {}
    for raw in paths:
        key = nfc(printable(os.path.abspath(raw)))
        if key in seen:
            if not same_file(seen[key], raw):
                entries.append(twin_entry(key, raw))
            continue
        seen[key] = raw
        entry = process_one(raw, key, out_dir, max_mb, inventory, taken)
        inventory[key] = entry
        entries.append(entry)
        save_inventory(out_dir, inventory)  # a cada arquivo, para uma interrupção não perder o que já foi gravado
    return entries


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument('arquivos', nargs='+', metavar='ARQUIVO', help='PDF, DOCX, XLSX, TXT ou MD')
    parser.add_argument('--saida', type=Path, required=True, help='pasta dos .txt e do inventario.json')
    parser.add_argument('--max-mb', type=float, default=200, help='tamanho máximo por arquivo (padrão: 200)')
    args = parser.parse_args(argv)
    try:
        entries = process(args.arquivos, args.saida, args.max_mb)
    except KeyboardInterrupt:
        print('interrompido: o inventário tem os arquivos já extraídos; rode de novo para completar',
              file=sys.stderr)
        return 1
    except ValueError as exc:
        print(f'erro: {exc}', file=sys.stderr)
        return 1
    except OSError as exc:
        print(f'erro: não foi possível ler ou gravar em {args.saida}: {exc.strerror or exc}', file=sys.stderr)
        return 1
    failed = [e for e in entries if e['erro']]
    for entry in failed:
        print(f"erro: {entry['nome']}: {entry['erro']}", file=sys.stderr)
    print(json.dumps({'diretorio': printable(str(args.saida)), 'inventario': printable(str(args.saida / INVENTORY)),
                      'processados': len(entries), 'com_erro': len(failed), 'arquivos': entries},
                     ensure_ascii=False, indent=2))
    return 1 if failed else 0


if __name__ == '__main__':
    sys.exit(main())
