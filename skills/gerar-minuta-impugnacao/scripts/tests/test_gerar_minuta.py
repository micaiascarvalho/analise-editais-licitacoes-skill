import io
import sys
import tempfile
import json
import unittest
from datetime import date
from unittest import mock
from contextlib import redirect_stderr, redirect_stdout
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml.ns import qn

import gerar_minuta
from gerar_minuta import Refusal, STYLES

# Dados fictícios.
MINUTA = '''---
tipo: impugnacao
---

**ILUSTRÍSSIMO SENHOR PREGOEIRO DA SECRETARIA MUNICIPAL DE EXEMPLÓPOLIS.**

**Referente:** Pregão Eletrônico nº 1/2026.

# DA TEMPESTIVIDADE

Texto com **negrito** e _itálico_.

# DOS PONTOS IMPUGNADOS

**III.1 – CAPITAL E PATRIMÔNIO CUMULATIVOS**

| # | Cláusula impugnada | Vício |
| :-- | :-- | :-- |
| III.1 | Item 8.6.3 | Cumulação vedada |

> “Trecho literal do edital.”

# DOS PEDIDOS

1. o **CONHECIMENTO** da impugnação;
1. o **PROVIMENTO** do pedido;

2. a retificação do item;
2. a republicação do edital;

<!-- centralizado -->
Teresina, 9 de setembro de 2026.
<!-- /centralizado -->

Fulano de Tal
Sócio Administrador
'''


class SemDadosReais:
    """Os testes não leem o assets/dados-servfaz.json de verdade (tem dado de pessoa física)."""
    def setUp(self):
        patcher = mock.patch.object(gerar_minuta, 'DADOS', Path('/nao/existe/dados.json'))
        patcher.start()
        self.addCleanup(patcher.stop)


class GerarMinutaTests(SemDadosReais, unittest.TestCase):
    def convert(self, text, name='minuta.md'):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        source = Path(self.tmp.name) / name
        source.write_text(text, encoding='utf-8')
        return source, Document(str(gerar_minuta.convert(source).docx))

    def test_conventions(self):
        source, doc = self.convert(MINUTA)
        self.assertEqual(source.with_suffix('.docx').name, 'minuta.docx')
        styles = [p.style.name for p in doc.paragraphs]
        texts = [p.text for p in doc.paragraphs]
        self.assertEqual(sum(1 for s in styles if s == STYLES['titulo']), 3)
        self.assertIn(STYLES['subtitulo'], styles)
        self.assertIn(STYLES['citacao'], styles)
        self.assertIn(STYLES['centro'], styles)
        self.assertEqual(doc.tables[0].cell(0, 1).text, 'Cláusula impugnada')
        self.assertFalse(any('**' in t or '---' in t for t in texts))
        body = next(p for p in doc.paragraphs if p.text.startswith('Texto com'))
        self.assertEqual([(r.text, r.bold, r.italic) for r in body.runs if r.bold or r.italic],
                         [('negrito', True, None), ('itálico', None, True)])
        self.assertTrue(doc.sections[0].header.paragraphs or doc.sections[0].header.tables)

    def test_lists_restart_and_close_with_period(self):
        _, doc = self.convert(MINUTA)
        alineas = [p for p in doc.paragraphs if p.style.name == STYLES['alinea']]
        lists = [p._p.pPr.find(qn('w:numPr')).find(qn('w:numId')).get(qn('w:val')) for p in alineas]
        self.assertEqual(len(set(lists)), 2)  # cada lista recomeça em a)
        self.assertTrue(alineas[1].text.endswith('.'), alineas[1].text)
        self.assertTrue(alineas[3].text.endswith('.'), alineas[3].text)

    def test_signature_block_has_no_space_between_lines(self):
        _, doc = self.convert(MINUTA)
        name = next(p for p in doc.paragraphs if p.text == 'Fulano de Tal')
        self.assertEqual(name.paragraph_format.space_after.pt, 0)

    def test_refuses_internal_references(self):
        for text in ('Fonte [EXT-LEI14133-ART5 · conferido].\n', 'Ver catálogo de fontes.\n'):
            with tempfile.TemporaryDirectory() as tmp:
                path = Path(tmp) / 'minuta.md'
                path.write_text(text, encoding='utf-8')
                with self.assertRaises(Refusal):
                    gerar_minuta.convert(path)
                self.assertFalse(path.with_suffix('.docx').exists())

    def test_catalog_of_services_is_allowed(self):
        self.convert('O Catálogo de Serviços do órgão não prevê o item.\n')

    def test_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / 'minuta.md'
            source.write_text(MINUTA, encoding='utf-8')
            target = Path(tmp) / 'saida.docx'
            with redirect_stdout(io.StringIO()):
                self.assertEqual(gerar_minuta.main([str(source), '--saida', str(target)]), 0)
            self.assertTrue(target.is_file())
            source.write_text('Ver catálogo de fontes.\n', encoding='utf-8')
            with redirect_stderr(io.StringIO()) as err:
                self.assertEqual(gerar_minuta.main([str(source)]), 1)
            self.assertIn('linha 1', err.getvalue())
            with redirect_stderr(io.StringIO()):
                self.assertEqual(gerar_minuta.main([str(Path(tmp) / 'nao-existe.md')]), 1)


# Markdown comum, como sai de um gerador de chat (dados fictícios).
COMUM = '''**ILUSTRÍSSIMO SENHOR PREGOEIRO.**

## IMPUGNAÇÃO AO EDITAL

---

### I. DA TEMPESTIVIDADE

Texto com *itálico* simples.

#### 1. Do primeiro ponto

> *"Trecho literal do edital."* [fonte]

Verifica-se que:
1. primeiro;
2. segundo;

### IV. DOS PEDIDOS

1. O **CONHECIMENTO** da impugnação;
2. No mérito, o **ACOLHIMENTO**, para que:
   - **a)** A exclusão do item 1;
   - **b)** A adequação do item 2;
3. A republicação do edital;

- solto um;
- solto dois;
'''


class MarkdownComumTests(SemDadosReais, unittest.TestCase):
    convert = GerarMinutaTests.convert

    def test_common_markdown(self):
        _, doc = self.convert(COMUM)
        rows = [(p.style.name, p.text) for p in doc.paragraphs]
        titles = [t for s, t in rows if s == STYLES['titulo']]
        self.assertEqual(titles, ['DA TEMPESTIVIDADE', 'DOS PEDIDOS'])  # "I." digitado sai: a numeração é automática
        self.assertIn((STYLES['centro'], 'IMPUGNAÇÃO AO EDITAL'), rows)
        self.assertIn((STYLES['subtitulo'], '1. Do primeiro ponto'), rows)
        self.assertTrue(all('---' not in t and '#' not in t and '*' not in t for _, t in rows))
        italic = next(p for p in doc.paragraphs if p.text.startswith('Texto com'))
        self.assertEqual([r.text for r in italic.runs if r.italic], ['itálico'])

    def test_nested_items_stay_in_one_list(self):
        _, doc = self.convert(COMUM)
        alineas = [p for p in doc.paragraphs if p.style.name == STYLES['alinea']]
        numbered = [p for p in alineas if p._p.pPr.find(qn('w:numPr')) is not None]
        merits = [p for p in numbered if p.text.startswith(('O CONHECIMENTO', 'No mérito', 'A republicação'))]
        ids = {p._p.pPr.find(qn('w:numPr')).find(qn('w:numId')).get(qn('w:val')) for p in merits}
        self.assertEqual(len(merits), 3)
        self.assertEqual(len(ids), 1)  # a republicação é c), não recomeça em a)
        subs = [p for p in alineas if p.text.startswith(('a) A exclusão', 'b) A adequação'))]
        self.assertEqual(len(subs), 2)
        self.assertTrue(all(p.paragraph_format.left_indent > p.style.paragraph_format.left_indent for p in subs))
        self.assertTrue(subs[1].text.endswith(';'))  # o ponto final vai ao último item da lista, não ao sub-item
        self.assertTrue(merits[2].text.endswith('.'))

    def test_loose_bullets_become_numbered(self):
        _, doc = self.convert(COMUM)
        loose = [p for p in doc.paragraphs if p.text.startswith('solto')]
        self.assertTrue(all(p._p.pPr.find(qn('w:numPr')) is not None for p in loose))
        self.assertTrue(loose[-1].text.endswith('.'))

    def test_refuses_unsupported_markup(self):
        cases = {'```\ncódigo\n```': '```',
                 '![logo](a.png)': '![logo]', 'Veja [o edital](http://x.br).': 'o edital', 'Linha<br>quebra': '<br>',
                 '##### Subtítulo fundo': '#####', 'Texto ~~riscado~~.': '~~'}
        for text, found in cases.items():
            with tempfile.TemporaryDirectory() as tmp:
                path = Path(tmp) / 'minuta.md'
                path.write_text(text + '\n', encoding='utf-8')
                with self.assertRaisesRegex(Refusal, found.replace('[', r'\[').replace('(', r'\(')):
                    gerar_minuta.convert(path)
                self.assertFalse(path.with_suffix('.docx').exists(), text)

    def test_numeric_references_pass_through(self):
        _, doc = self.convert('Conforme o edital [1] e o TR [2, 3].\n')
        self.assertEqual(doc.paragraphs[-1].text, 'Conforme o edital [1] e o TR [2, 3].')

    def test_legal_text_with_brackets_and_comparisons_is_allowed(self):
        self.convert('Índice [sic] inferior a 1,0 e valor < 10 ou > 5; art. 67, § 2º.\n')


DADOS_TESTE = {'cnpj': '00.000.000/0001-00', 'endereco': 'Rua Exemplo, n.º 1, Exemplópolis – UF', 'nome': 'Fulana de Tal',
               'rg': '0.000.000-SSP/UF', 'cpf': '000.000.000-00'}
TEXTO_CAMPOS = ('Inscrita no CNPJ nº [Inserir CNPJ], com sede em [Inserir Endereço Completo].\n\n'
                'Representante: [Inserir Nome], RG [Inserir RG], CPF: [Inserir CPF].\n\n'
                'Teresina/PI, [Data do Protocolo].\n')


class DadosPadraoTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.dir = Path(self.tmp.name)
        self.source = self.dir / 'minuta.md'
        self.dados = self.dir / 'dados.json'
        self.dados.write_text(json.dumps(DADOS_TESTE), encoding='utf-8')

    def texts(self, text, **kwargs):
        self.source.write_text(text, encoding='utf-8')
        result = gerar_minuta.convert(self.source, dados=self.dados, **kwargs)
        self.result = result
        return [p.text for p in Document(str(result.docx)).paragraphs if p.text]

    def test_fields_filled_from_data_and_date(self):
        texts = self.texts(TEXTO_CAMPOS, when=date(2026, 3, 9))
        self.assertEqual(texts, ['Inscrita no CNPJ nº 00.000.000/0001-00, com sede em Rua Exemplo, n.º 1, Exemplópolis – UF.',
                                 'Representante: Fulana de Tal, RG 0.000.000-SSP/UF, CPF: 000.000.000-00.',
                                 'Teresina/PI, 9 de março de 2026.'])

    def test_missing_fields_do_not_block_and_are_highlighted(self):
        texts = self.texts(TEXTO_CAMPOS)  # sem data
        self.assertEqual(texts[-1], 'Teresina/PI, [Data do Protocolo].')
        self.assertEqual(self.result.pendentes, [(5, '[Data do Protocolo]')])
        doc = Document(str(self.result.docx))
        marked = [r.text for p in doc.paragraphs for r in p.runs if r.font.highlight_color == WD_COLOR_INDEX.YELLOW]
        self.assertEqual(marked, ['[Data do Protocolo]'])

    def test_unknown_field_and_template_marker_are_pending(self):
        self.texts('Telefone [Inserir Telefone] e {{CAMPO}}.\n')
        self.assertEqual([m for _, m in self.result.pendentes], ['[Inserir Telefone]', '{{CAMPO}}'])

    def test_missing_data_file_leaves_everything_pending(self):
        self.source.write_text('CNPJ [Inserir CNPJ].\n', encoding='utf-8')
        result = gerar_minuta.convert(self.source, dados=self.dir / 'nao-existe.json')
        self.assertEqual(result.pendentes, [(1, '[Inserir CNPJ]')])

    def test_cli_lists_pending(self):
        self.source.write_text('Teresina/PI, [Data do Protocolo].\n', encoding='utf-8')
        with redirect_stdout(io.StringIO()) as out:
            self.assertEqual(gerar_minuta.main([str(self.source)]), 0)
        self.assertIn('1 campo(s) em branco', out.getvalue())
        self.assertIn('linha 1 do .md: [Data do Protocolo]', out.getvalue())

    def test_invalid_data_file(self):
        self.dados.write_text('{"cnpj": 1}', encoding='utf-8')
        with self.assertRaisesRegex(Refusal, 'valores em texto'):
            self.texts('Texto.\n')

    def test_example_file_has_every_field(self):
        example = json.loads((Path(gerar_minuta.__file__).resolve().parents[1] / 'assets' /
                              'dados-servfaz.exemplo.json').read_text(encoding='utf-8'))
        self.assertTrue(set(DADOS_TESTE) <= set(example))

    def test_cli_date(self):
        self.source.write_text('Teresina/PI, [Data do Protocolo].\n', encoding='utf-8')
        with redirect_stderr(io.StringIO()) as err:
            self.assertEqual(gerar_minuta.main([str(self.source), '--data', '2026-13-01']), 1)
        self.assertIn('data inválida', err.getvalue())


class TimbreTests(unittest.TestCase):
    def test_footer_line_is_above_the_text(self):
        """A linha laranja do rodapé não pode cruzar o texto "servfazpi | www.servfaz.com.br"."""
        footer = Document(str(gerar_minuta.TIMBRE)).sections[0].footer._element
        offset = lambda name: int(footer.xpath(f'.//wp:anchor[wp:docPr/@name="{name}"]/wp:positionV/wp:posOffset')[0].text)
        line_bottom = offset('Conector reto 9') + int(footer.xpath('.//wp:anchor[wp:docPr/@name="Conector reto 9"]/wp:extent')[0].get('cy'))
        text_top = offset('Caixa de Texto 2') + 45720  # recuo interno superior da caixa
        self.assertLess(line_bottom, text_top)


if __name__ == '__main__':
    unittest.main()
