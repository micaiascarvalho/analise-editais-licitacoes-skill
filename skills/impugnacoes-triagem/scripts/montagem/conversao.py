"""Conversão da peça montada (.md) em .docx sobre o timbre, pelas convenções de assets/LEIA-ME.md: "# " é seção com
numeração romana, "> " é citação recuada, "1. " é alínea a), b), tabela "|" é o quadro, parágrafo todo em negrito é
subtítulo e <!-- centralizado --> centraliza o fecho. Cada linha não vazia é um parágrafo; linha seguida de outra sem
linha em branco entre elas sai sem espaço depois (bloco de assinatura)."""
import re
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

from triagem.comum import Refusal, nfc

from .modelo import check_clean

# Nomes em ASCII: o nome do estilo vira o styleId no XML, e leitores de .docx tropeçam em acento ali.
STYLES = {'corpo': 'Peca Corpo', 'titulo': 'Peca Titulo', 'subtitulo': 'Peca Subtitulo', 'citacao': 'Peca Citacao',
          'alinea': 'Peca Alinea', 'centro': 'Peca Centro', 'tabela': 'Peca Tabela'}
ROMAN_NUMBERING, ALINEA_NUMBERING = 'PecaSecao', 'PecaAlinea'
TIMBRE = Path(__file__).resolve().parents[2] / 'assets' / 'timbre-servfaz.docx'
ITEM = re.compile(r'^\d+\.\s+(.*)$')
BOLD_LINE = re.compile(r'^\*\*[^*]+\*\*$')
INLINE = re.compile(r'(\*\*.+?\*\*|(?<![\w*])_(?!_).+?(?<!_)_(?!\w))')
CENTER_ON, CENTER_OFF = '<!-- centralizado -->', '<!-- /centralizado -->'


def add_runs(paragraph, text):
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith('**') and part.endswith('**') and len(part) > 4:
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith('_') and part.endswith('_') and len(part) > 2:
            paragraph.add_run(part[1:-1]).italic = True
        else:
            paragraph.add_run(part)
    return paragraph


def new_list(doc):
    """Nova numeração a), b) que recomeça em a): um w:num por lista, com startOverride."""
    numbering = doc.part.numbering_part.element
    abstract_id = None
    for abstract in numbering.findall(qn('w:abstractNum')):
        name = abstract.find(qn('w:name'))
        if name is not None and name.get(qn('w:val')) == ALINEA_NUMBERING:
            abstract_id = abstract.get(qn('w:abstractNumId'))
    if abstract_id is None:
        raise Refusal(f'o timbre não tem a numeração {ALINEA_NUMBERING}: gere-o com criar_timbre.py.')
    num_id = max([int(n.get(qn('w:numId'))) for n in numbering.findall(qn('w:num'))], default=0) + 1
    num = OxmlElement('w:num')
    num.set(qn('w:numId'), str(num_id))
    ref = OxmlElement('w:abstractNumId')
    ref.set(qn('w:val'), abstract_id)
    override = OxmlElement('w:lvlOverride')
    override.set(qn('w:ilvl'), '0')
    start = OxmlElement('w:startOverride')
    start.set(qn('w:val'), '1')
    override.append(start)
    num.extend([ref, override])
    numbering.append(num)
    return num_id


def number_paragraph(paragraph, num_id):
    numpr = OxmlElement('w:numPr')
    for tag, value in (('w:ilvl', '0'), ('w:numId', str(num_id))):
        node = OxmlElement(tag)
        node.set(qn('w:val'), value)
        numpr.append(node)
    paragraph._p.get_or_add_pPr().append(numpr)


def add_table(doc, lines):
    rows = [[cell.strip() for cell in line.strip().strip('|').split('|')] for line in lines]
    rows = [row for row in rows if not all(re.fullmatch(r':?-+:?', cell) for cell in row)]
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    borders = OxmlElement('w:tblBorders')
    for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        edge = OxmlElement(f'w:{side}')
        for key, value in (('val', 'single'), ('sz', '4'), ('space', '0'), ('color', '000000')):
            edge.set(qn(f'w:{key}'), value)
        borders.append(edge)
    table._tbl.tblPr.append(borders)
    for r, row in enumerate(rows):
        for c, text in enumerate(row):
            paragraph = table.cell(r, c).paragraphs[0]
            paragraph.style = STYLES['tabela']
            add_runs(paragraph, f'**{text}**' if r == 0 and text else text)
    header = OxmlElement('w:tblHeader')
    table.rows[0]._tr.get_or_add_trPr().append(header)
    doc.add_paragraph(style=STYLES['corpo'])


def build(doc, text):
    lines = text.split('\n')
    centered, i = False, 0
    while i < len(lines):
        line = lines[i].strip()
        following = lines[i + 1].strip() if i + 1 < len(lines) else ''
        i += 1
        if not line:
            continue
        if line in (CENTER_ON, CENTER_OFF):
            centered = line == CENTER_ON
            continue
        if line.startswith('|'):
            block = [line]
            while i < len(lines) and lines[i].strip().startswith('|'):
                block.append(lines[i].strip())
                i += 1
            add_table(doc, block)
            continue
        if ITEM.match(line):
            items = [ITEM.match(line)[1]]
            while i < len(lines) and ITEM.match(lines[i].strip()):
                items.append(ITEM.match(lines[i].strip())[1])
                i += 1
            items[-1] = re.sub(r';\s*$', '.', items[-1])
            num_id = new_list(doc)
            for item in items:
                number_paragraph(add_runs(doc.add_paragraph(style=STYLES['alinea']), item), num_id)
            continue
        if line.startswith('# '):
            add_runs(doc.add_paragraph(style=STYLES['titulo']), line[2:].strip())
            continue
        if line.startswith('>'):
            add_runs(doc.add_paragraph(style=STYLES['citacao']), line[1:].strip())
            continue
        style = STYLES['centro'] if centered else STYLES['subtitulo'] if BOLD_LINE.match(line) else STYLES['corpo']
        paragraph = add_runs(doc.add_paragraph(style=style), line)
        if following and following not in (CENTER_ON, CENTER_OFF):
            paragraph.paragraph_format.space_after = Pt(0)


def convert(source, target=None, timbre=TIMBRE):
    source = Path(source)
    text = nfc(source.read_text(encoding='utf-8'))
    check_clean(text, source.name)
    if not Path(timbre).is_file():
        raise Refusal(f'timbre não encontrado ({timbre}): gere-o com criar_timbre.py.')
    doc = Document(str(timbre))
    build(doc, text)
    target = Path(target) if target else source.with_suffix('.docx')
    doc.save(str(target))
    return target
