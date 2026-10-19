import io
import json
import os
import sys
import tempfile
import tracemalloc
import unicodedata
import unittest
import warnings
import zipfile
from contextlib import redirect_stderr, redirect_stdout
from pathlib import Path
from unittest import mock

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import extrair_texto
from extrair_texto import NO_TEXT, extract_docx, extract_pdf, extract_xlsx, file_sha256, process

W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
MC = 'http://schemas.openxmlformats.org/markup-compatibility/2006'


def make_pdf(path, pages):
    """PDF mínimo montado à mão: uma página por texto (lista = uma linha por item; None = página sem conteúdo)."""
    objects = ['<< /Type /Catalog /Pages 2 0 R >>', None,
               '<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>']
    kids = []
    for text in pages:
        if isinstance(text, list):
            text = ') Tj 0 -14 Td ('.join(text)
        content = f'BT /F1 12 Tf 72 720 Td ({text}) Tj ET' if text else ''
        objects.append(f'<< /Length {len(content)} >>\nstream\n{content}\nendstream')
        objects.append(f'<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] '
                       f'/Resources << /Font << /F1 3 0 R >> >> /Contents {len(objects)} 0 R >>')
        kids.append(f'{len(objects)} 0 R')
    objects[1] = f"<< /Type /Pages /Kids [{' '.join(kids)}] /Count {len(kids)} >>"
    out, offsets = b'%PDF-1.4\n', []
    for n, body in enumerate(objects, 1):
        offsets.append(len(out))
        out += f'{n} 0 obj\n{body}\nendobj\n'.encode('ascii')
    xref = len(out)
    out += f'xref\n0 {len(objects) + 1}\n0000000000 65535 f \n'.encode('ascii')
    out += ''.join(f'{offset:010d} 00000 n \n' for offset in offsets).encode('ascii')
    out += f'trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n'.encode('ascii')
    path.write_bytes(out)
    return path


def encrypt_pdf(source, target, user_password):
    from pypdf import PdfReader, PdfWriter
    writer = PdfWriter(clone_from=PdfReader(source))
    writer.encrypt(user_password, owner_password='dono', algorithm='RC4-128')
    with open(target, 'wb') as fh:
        writer.write(fh)
    return target


def paragraph(text):
    return f'<w:p><w:r><w:t xml:space="preserve">{text}</w:t></w:r></w:p>'


def make_docx(path, body, extra=None):
    document = (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                f'<w:document xmlns:w="{W}" xmlns:mc="{MC}"><w:body>{body}</w:body></w:document>')
    with zipfile.ZipFile(path, 'w', zipfile.ZIP_DEFLATED) as zf:
        zf.writestr('[Content_Types].xml', '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"/>')
        zf.writestr('word/document.xml', document)
        for name, data in (extra or {}).items():
            zf.writestr(name, data)
    return path


def make_xlsx(path):
    from openpyxl import Workbook
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Custos'
    sheet['A1'], sheet['C1'] = 'Posto', 'Valor mensal'
    sheet['A2'], sheet['C2'] = 'Servente fictício', 1500.5
    sheet['B4'] = 'linha\ncom quebra'
    workbook.create_sheet('Vazia')
    hidden = workbook.create_sheet('Oculta')
    hidden['A1'] = 'interno'
    hidden.sheet_state = 'hidden'
    workbook.save(path)
    return path


def run(*args):
    out, err = io.StringIO(), io.StringIO()
    with redirect_stdout(out), redirect_stderr(err):
        code = extrair_texto.main([str(a) for a in args])
    return code, out.getvalue(), err.getvalue()


class ExtrairTextoTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.base = Path(self.tmp.name)
        self.docs = self.base / 'docs'
        self.docs.mkdir()
        self.out = self.base / 'texto'

    def tearDown(self):
        self.tmp.cleanup()

    def inventory(self):
        return json.loads((self.out / 'inventario.json').read_text(encoding='utf-8'))

    def output(self, entry):
        return (self.out / entry['saida']).read_text(encoding='utf-8')

    def test_pdf_marks_each_page_and_blank_pages(self):
        pdf = make_pdf(self.docs / 'Edital PE 1.2026.pdf', ['Objeto da licita\\347\\343o fict\\355cia', None])
        text, pages, blank = extract_pdf(pdf)
        self.assertEqual((pages, blank), (2, [2]))
        self.assertIn('=== página PDF 1 ===\nObjeto da licitação fictícia', text)
        self.assertIn(f'=== página PDF 2 ===\n{NO_TEXT}', text)
        [entry] = process([pdf], self.out)
        self.assertIsNone(entry['erro'])
        self.assertEqual(entry['saida'], 'edital-pe-1-2026.pdf.txt')
        self.assertEqual((entry['tipo'], entry['paginas'], entry['paginas_sem_texto']), ('pdf', 2, [2]))
        self.assertEqual(entry['caracteres'], len('Objeto da licitação fictícia'))
        self.assertEqual(entry['sha256'], file_sha256(pdf))
        head = self.output(entry).splitlines()[:5]
        self.assertEqual(head, ['# Fonte: Edital PE 1.2026.pdf', f"# SHA-256: {entry['sha256']}", '# Tipo: pdf',
                                '# Páginas: 2', '# Páginas sem texto: 2'])

    def test_page_that_breaks_the_extractor_is_marked_and_the_rest_is_kept(self):
        from pypdf import PageObject
        pdf = make_pdf(self.docs / 'malformado.pdf', ['Pagina boa', 'Outra pagina'])
        with mock.patch.object(PageObject, 'extract_text', side_effect=[KeyError('/Font'), 'Outra pagina']):
            [entry] = process([pdf], self.out)
        self.assertIsNone(entry['erro'])
        self.assertEqual(entry['paginas_sem_texto'], [1])
        self.assertIn('=== página PDF 1 ===\n[falha ao extrair esta página (KeyError): pode exigir leitura visual]',
                      self.output(entry))
        self.assertEqual(entry['caracteres'], len('Outra pagina'))

    def test_pdf_with_empty_user_password_opens_and_real_password_fails(self):
        plain = make_pdf(self.docs / 'base.pdf', ['Texto ficticio'])
        open_pdf = encrypt_pdf(plain, self.docs / 'aberto.pdf', '')
        locked = encrypt_pdf(plain, self.docs / 'trancado.pdf', 'segredo')
        self.assertIn('Texto ficticio', extract_pdf(open_pdf)[0])
        opened, failed = process([open_pdf, locked], self.out)
        self.assertIsNone(opened['erro'])
        self.assertIn('senha', failed['erro'])
        self.assertIsNone(failed['saida'])

    def test_missing_pypdf_points_to_requirements_and_native_reader(self):
        pdf = make_pdf(self.docs / 'edital.pdf', ['x'])
        with mock.patch.dict(sys.modules, {'pypdf': None}):
            [entry] = process([pdf], self.out)
        self.assertIn('pip install -r scripts/requirements.txt', entry['erro'])
        self.assertIn('Read com pages', entry['erro'])

    def test_docx_numbers_non_empty_paragraphs_including_table_cells(self):
        textbox = ('<w:p><w:r><w:t>Antes da caixa</w:t></w:r><mc:AlternateContent><mc:Choice Requires="wps">'
                   f'<w:drawing><w:txbxContent>{paragraph("Na caixa")}</w:txbxContent></w:drawing></mc:Choice>'
                   f'<mc:Fallback><w:pict><w:txbxContent>{paragraph("Na caixa")}</w:txbxContent></w:pict>'
                   '</mc:Fallback></mc:AlternateContent></w:p>')
        body = (paragraph('Cláusula primeira fictícia') + '<w:p/>' + paragraph('   ')
                + '<w:tbl><w:tr><w:tc>' + paragraph('Célula A') + '</w:tc><w:tc>' + paragraph('Célula B')
                + '</w:tc></w:tr></w:tbl>' + '<w:p><w:r><w:t>Dividido</w:t><w:tab/><w:t>em runs</w:t></w:r></w:p>'
                + textbox)
        docx = make_docx(self.docs / 'Minuta.docx', body)
        self.assertEqual(extract_docx(docx).splitlines(), [
            '[¶1] Cláusula primeira fictícia', '[¶2] Célula A', '[¶3] Célula B', '[¶4] Dividido em runs',
            '[¶5] Antes da caixa', '[¶6] Na caixa'])
        [entry] = process([docx], self.out)
        text = self.output(entry)
        self.assertIn('# Tipo: docx\n# Paginação: não verificada (DOCX)\n# Parágrafos: [¶N] numera só os não vazios', text)
        self.assertEqual(entry['saida'], 'minuta.docx.txt')
        self.assertIsNone(entry['paginas'])

    def test_xlsx_lists_sheets_and_non_empty_cells(self):
        xlsx = make_xlsx(self.docs / 'Planilha de Custos.xlsx')
        text = extract_xlsx(xlsx)
        self.assertIn('=== aba: Custos ===\nL1: A=Posto | C=Valor mensal\nL2: A=Servente fictício | C=1500.5\n'
                      'L4: B=linha com quebra', text)
        self.assertIn('=== aba: Vazia ===\n[aba sem células preenchidas]', text)
        self.assertIn('=== aba: Oculta ===\n[aba oculta]\nL1: A=interno', text)
        [entry] = process([xlsx], self.out)
        self.assertIn('# Valores: em cache da planilha; fórmulas não recalculadas', self.output(entry))
        self.assertEqual(entry['saida'], 'planilha-de-custos.xlsx.txt')

    def test_text_file_falls_back_to_cp1252(self):
        (self.docs / 'cct.txt').write_bytes('Cláusula fictícia'.encode('cp1252'))
        (self.docs / 'notas.md').write_text('# Notas\n\nTexto fictício\n', encoding='utf-8')
        cct, notes = process([self.docs / 'cct.txt', self.docs / 'notas.md'], self.out)
        self.assertIn('# Codificação: cp1252\n\n=== texto ===\nCláusula fictícia\n', self.output(cct))
        self.assertIn('=== texto ===\n# Notas\n\nTexto fictício\n', self.output(notes))
        self.assertEqual(notes['tipo'], 'texto')

    def test_nfd_name_is_opened_as_given_and_recorded_in_nfc(self):
        folder = self.docs / unicodedata.normalize('NFD', 'Prefeitura São Luís')
        folder.mkdir()
        name = 'ANEXO I - TERMO DE REFERÊNCIA Nº 1.pdf'
        pdf = make_pdf(folder / unicodedata.normalize('NFD', name), ['Termo ficticio'])
        [entry] = process([pdf], self.out)
        self.assertIsNone(entry['erro'])
        self.assertEqual(entry['nome'], name)
        self.assertEqual(entry['arquivo'], unicodedata.normalize('NFC', str(pdf)))
        self.assertEqual(entry['saida'], 'anexo-i-termo-de-referencia-no-1.pdf.txt')
        self.assertIn(f'# Fonte: {name}\n', self.output(entry))
        # O mesmo arquivo digitado em NFC é achado e não duplica o inventário.
        [again] = process([unicodedata.normalize('NFC', str(pdf))], self.out)
        self.assertIsNone(again['erro'])
        self.assertEqual(len(self.inventory()), 1)
        self.assertEqual(extract_pdf(self.inventory()[0]['arquivo'])[1], 1)

    def test_name_collision_gets_suffix_and_stays_stable(self):
        for folder in ('a', 'b', 'c'):
            (self.docs / folder).mkdir()
        first = make_pdf(self.docs / 'a' / 'Edital.pdf', ['A'])
        second = make_pdf(self.docs / 'b' / 'EDITAL.pdf', ['B'])
        third = make_pdf(self.docs / 'c' / 'Édital.pdf', ['C'])
        names = [e['saida'] for e in process([first, second, third], self.out)]
        self.assertEqual(names, ['edital.pdf.txt', 'edital-2.pdf.txt', 'edital-3.pdf.txt'])
        self.assertEqual([e['saida'] for e in process([third, first], self.out)], ['edital-3.pdf.txt', 'edital.pdf.txt'])
        self.assertIn('=== página PDF 1 ===\nB', (self.out / 'edital-2.pdf.txt').read_text(encoding='utf-8'))

    def test_unsupported_format_fails_with_code_1_and_others_are_written(self):
        (self.docs / 'antigo.doc').write_bytes(b'\xd0\xcf\x11\xe0 ficticio')
        (self.docs / 'nota.txt').write_text('Nota fictícia', encoding='utf-8')
        code, out, err = run(self.docs / 'antigo.doc', self.docs / 'nota.txt', '--saida', self.out)
        self.assertEqual(code, 1)
        self.assertIn('antigo.doc: formato não suportado (.doc)', err)
        summary = json.loads(out)
        self.assertEqual((summary['processados'], summary['com_erro']), (2, 1))
        bad, good = summary['arquivos']
        self.assertIsNone(bad['tipo'])
        self.assertIsNone(bad['saida'])
        self.assertTrue((self.out / good['saida']).exists())
        self.assertEqual([e['nome'] for e in self.inventory()], ['antigo.doc', 'nota.txt'])

    def test_all_ok_exits_0_and_missing_file_is_an_error(self):
        (self.docs / 'nota.txt').write_text('ok', encoding='utf-8')
        code, out, _ = run(self.docs / 'nota.txt', '--saida', self.out)
        self.assertEqual(code, 0)
        self.assertEqual(json.loads(out)['com_erro'], 0)
        code, _, err = run(self.docs / 'sumiu.pdf', '--saida', self.out)
        self.assertEqual(code, 1)
        self.assertIn('sumiu.pdf: arquivo não encontrado', err)

    def test_file_above_limit_is_refused(self):
        big = self.docs / 'grande.txt'
        big.write_bytes(b'x' * 2048)
        [entry] = process([big], self.out, max_mb=0.001)
        self.assertIn('maior que o limite', entry['erro'])
        self.assertIsNone(entry['sha256'])

    def test_zip_bomb_is_refused(self):
        bomb = make_docx(self.docs / 'bomba.docx', paragraph('x'), {'word/media/zeros.bin': b'\0' * (2 * 1024 * 1024)})
        self.assertLess(bomb.stat().st_size, 100 * 1024)
        [entry] = process([bomb], self.out, max_mb=0.2)
        self.assertIn('possível zip bomb', entry['erro'])
        self.assertIsNone(entry['saida'])

    def test_corrupted_office_file_is_an_error(self):
        (self.docs / 'falso.xlsx').write_bytes(b'nao e zip')
        [entry] = process([self.docs / 'falso.xlsx'], self.out)
        self.assertIn('não é um documento Office', entry['erro'])

    def test_rerun_updates_inventory_without_duplicates(self):
        note = self.docs / 'nota.txt'
        note.write_text('versão 1', encoding='utf-8')
        other = make_pdf(self.docs / 'edital.pdf', ['Edital ficticio'])
        process([note, other], self.out)
        first = {e['nome']: e for e in self.inventory()}
        note.write_text('versão 2, retificada', encoding='utf-8')
        process([note, note], self.out)
        items = self.inventory()
        self.assertEqual([e['nome'] for e in items], ['edital.pdf', 'nota.txt'])
        self.assertEqual(items, sorted(items, key=lambda e: e['arquivo']))
        updated = {e['nome']: e for e in items}
        self.assertNotEqual(updated['nota.txt']['sha256'], first['nota.txt']['sha256'])
        self.assertEqual(updated['edital.pdf'], first['edital.pdf'])
        self.assertIn('versão 2, retificada', self.output(updated['nota.txt']))
        self.assertEqual(sorted(p.name for p in self.out.iterdir()), ['edital.pdf.txt', 'inventario.json', 'nota.txt.txt'])

    def test_file_that_now_fails_loses_its_old_text(self):
        note = self.docs / 'nota.txt'
        note.write_text('x' * 4096, encoding='utf-8')
        [entry] = process([note], self.out)
        self.assertTrue((self.out / entry['saida']).exists())
        [entry] = process([note], self.out, max_mb=0.001)
        self.assertIsNotNone(entry['erro'])
        self.assertFalse((self.out / 'nota.txt.txt').exists())

    def test_invalid_inventory_stops_without_writing(self):
        self.out.mkdir()
        (self.out / 'inventario.json').write_text('{quebrado', encoding='utf-8')
        (self.docs / 'nota.txt').write_text('ok', encoding='utf-8')
        code, _, err = run(self.docs / 'nota.txt', '--saida', self.out)
        self.assertEqual(code, 1)
        self.assertIn('inventario.json inválido', err)
        self.assertFalse((self.out / 'nota.txt.txt').exists())

    def test_aes_pdf_without_cryptography_is_a_file_error_not_blank_pages(self):
        from pypdf import PageObject
        from pypdf.errors import DependencyError
        pdf = make_pdf(self.docs / 'cifrado.pdf', ['Texto ficticio', 'Outro texto'])
        missing = DependencyError('cryptography>=3.1 is required for AES algorithm')
        with mock.patch.object(PageObject, 'extract_text', side_effect=missing):  # AES-128: abre, a página não
            [entry] = process([pdf], self.out)
        self.assertIn('pacote cryptography', entry['erro'])
        self.assertIn('Read com pages', entry['erro'])
        self.assertEqual((entry['saida'], entry['paginas_sem_texto']), (None, None))
        with mock.patch('pypdf.PdfReader', side_effect=missing):  # AES-256: falha já ao abrir
            code, _, err = run(pdf, '--saida', self.out)
        self.assertEqual(code, 1)
        self.assertIn('cifrado.pdf: PDF cifrado com AES exige o pacote cryptography', err)

    def test_aes256_pdf_with_empty_password_is_extracted(self):
        try:
            import cryptography  # noqa: F401
        except ImportError:
            self.skipTest('cryptography não instalado')
        from pypdf import PdfReader, PdfWriter
        plain = make_pdf(self.docs / 'claro.pdf', ['Texto ficticio cifrado'])
        writer = PdfWriter(clone_from=PdfReader(plain))
        writer.encrypt(user_password='', owner_password='dono', algorithm='AES-256')
        cipher = self.docs / 'aes.pdf'
        with cipher.open('wb') as handle:
            writer.write(handle)
        [entry] = process([cipher], self.out)
        self.assertIsNone(entry['erro'])
        self.assertIn('Texto ficticio cifrado', (self.out / entry['saida']).read_text(encoding='utf-8'))

    def test_inventory_output_must_be_a_txt_inside_the_folder(self):
        note = self.docs / 'nota.txt'
        note.write_text('Nota fictícia', encoding='utf-8')
        for saida in ('../fora.txt', str(self.base / 'absoluto.txt'), str(note), 5, ['x'], 'inventario.json'):
            with self.subTest(saida=saida):
                process([note], self.out)
                items = self.inventory()
                items[0]['saida'] = saida
                (self.out / 'inventario.json').write_text(json.dumps(items), encoding='utf-8')
                code, _, err = run(note, '--saida', self.out)
                self.assertEqual(code, 1)
                self.assertIn("inventario.json inválido ('saida'", err)
                self.assertEqual(sorted(p.name for p in self.base.iterdir()), ['docs', 'texto'])
                self.assertEqual(note.read_text(encoding='utf-8'), 'Nota fictícia')
                (self.out / 'inventario.json').unlink()
        process([note], self.out)
        items = self.inventory()
        items[0]['saida'] = 'renomeado.txt'  # renomear dentro da pasta continua valendo
        (self.out / 'inventario.json').write_text(json.dumps(items), encoding='utf-8')
        [entry] = process([note], self.out)
        self.assertEqual(entry['saida'], 'renomeado.txt')
        self.assertIn('Nota fictícia', self.output(entry))

    def test_non_utf8_file_name_is_an_error_and_the_others_are_saved(self):
        bad = self.docs / os.fsdecode(b'edital-s\xe3o.txt')
        try:
            bad.write_text('Texto fictício', encoding='utf-8')
        except (OSError, UnicodeEncodeError):
            self.skipTest('o sistema de arquivos não aceita nome fora do UTF-8')
        good = self.docs / 'nota.txt'
        good.write_text('Nota fictícia', encoding='utf-8')
        code, out, err = run(good, bad, '--saida', self.out)
        self.assertEqual(code, 1)
        summary = json.loads(out)
        self.assertEqual((summary['processados'], summary['com_erro']), (2, 1))
        self.assertIn('fora do UTF-8', err)
        self.assertEqual([e['nome'] for e in self.inventory()], ['edital-s�o.txt', 'nota.txt'])
        self.assertTrue((self.out / 'nota.txt.txt').exists())

    def test_docx_is_read_as_a_stream_without_building_the_whole_tree(self):
        body = paragraph('Início fictício') + paragraph(' ') * 50000 + paragraph('Fim fictício')
        docx = make_docx(self.docs / 'longo.docx', body)
        with zipfile.ZipFile(docx) as zf:
            xml_size = zf.getinfo('word/document.xml').file_size
        tracemalloc.start()
        try:
            text = extract_docx(docx)
            peak = tracemalloc.get_traced_memory()[1]
        finally:
            tracemalloc.stop()
        self.assertEqual(text, '[¶1] Início fictício\n[¶2] Fim fictício\n')
        self.assertLess(peak, xml_size // 2)

    def test_interrupted_run_keeps_inventory_of_finished_files(self):
        notes = [self.docs / f'nota{n}.txt' for n in range(3)]
        for n, note in enumerate(notes):
            note.write_text(f'Nota fictícia {n}', encoding='utf-8')
        first = extrair_texto.read_text(notes[0])
        with mock.patch.object(extrair_texto, 'read_text', side_effect=[first, KeyboardInterrupt]):
            try:
                code, _, err = run(*notes, '--saida', self.out)
            except KeyboardInterrupt:
                self.fail('a interrupção escapou do main')
        self.assertEqual(code, 1)
        self.assertIn('interrompido', err)
        self.assertEqual([e['nome'] for e in self.inventory()], ['nota0.txt'])
        self.assertEqual(sorted(p.name for p in self.out.iterdir()), ['inventario.json', 'nota0.txt.txt'])

    def test_common_failures_have_clear_messages(self):
        (self.docs / 'vazio.pdf').write_bytes(b'')
        (self.docs / 'pagina.pdf').write_bytes(b'<!doctype html><html><body>Erro ficticio</body></html>')
        broken = self.docs / 'quebrado.docx'
        with zipfile.ZipFile(broken, 'w') as zf:
            zf.writestr('word/document.xml', f'<w:document xmlns:w="{W}"><w:body><w:p>')
        empty, html, xml = process([self.docs / 'vazio.pdf', self.docs / 'pagina.pdf', broken], self.out)
        self.assertTrue(empty['erro'].startswith('arquivo vazio'))
        self.assertTrue(html['erro'].startswith('não é um PDF válido ou está corrompido'))
        self.assertIn('Read com pages', html['erro'])
        self.assertTrue(xml['erro'].startswith('XML interno do documento inválido'))

    def test_output_folder_problem_is_reported_in_portuguese(self):
        (self.docs / 'nota.txt').write_text('ok', encoding='utf-8')
        code, out, err = run(self.docs / 'nota.txt', '--saida', self.docs / 'nota.txt')
        self.assertEqual((code, out), (1, ''))
        self.assertIn('erro: não foi possível ler ou gravar em', err)
        self.assertNotIn('[Errno', err)

    def test_max_mb_must_be_positive_and_finite(self):
        (self.docs / 'nota.txt').write_text('ok', encoding='utf-8')
        for value in ('nan', 'inf', '-1', '0'):
            with self.subTest(value=value):
                code, _, err = run(self.docs / 'nota.txt', '--saida', self.out, '--max-mb', value)
                self.assertEqual(code, 1)
                self.assertIn('erro: --max-mb deve ser um número positivo e finito', err)
        self.assertFalse(self.out.exists())
        with self.assertRaises(ValueError):
            process([self.docs / 'nota.txt'], self.out, max_mb=float('nan'))

    def test_xlsx_out_of_range_date_is_flagged_without_raw_warnings(self):
        from openpyxl import Workbook
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = 'Datas'
        sheet['A1'], sheet['B1'] = 'Vigência fictícia', 99999999
        sheet['A2'], sheet['B2'] = 'Início fictício', 45000
        sheet['B1'].number_format = sheet['B2'].number_format = 'dd/mm/yyyy'
        workbook.save(self.docs / 'datas.xlsx')
        with warnings.catch_warnings(record=True) as caught:
            warnings.simplefilter('always')
            text = extract_xlsx(self.docs / 'datas.xlsx')
        self.assertEqual(caught, [])
        self.assertIn('=== aba: Datas ===\n[aviso: data fora do intervalo do Excel em B1; aparece como #VALUE!', text)
        self.assertIn('L2: A=Início fictício | B=2023-03-15', text)

    def test_pdf_line_that_looks_like_a_marker_is_neutralized(self):
        forged = ['Item 1 ficticio', '=== p\\341gina PDF 9 ===', '[p\\341gina sem texto extra\\355vel]', 'Item 2']
        pdf = make_pdf(self.docs / 'forjado.pdf', [forged, ['=== p\\341gina PDF 3 ===', 'Dois']])
        text, pages, blank = extract_pdf(pdf)
        self.assertEqual((pages, blank), (2, []))
        markers = [line for line in text.splitlines() if line.startswith(('===', '['))]
        self.assertEqual(markers, ['=== página PDF 1 ===', '=== página PDF 2 ==='])
        self.assertIn('\n === página PDF 9 ===\n', text)
        self.assertIn('=== página PDF 2 ===\n === página PDF 3 ===\n', text)

    def test_utf16_and_utf32_text_files_are_decoded_by_bom(self):
        text = 'Cláusula 5ª – “reajuste” de 3,5% — fictícia'
        (self.docs / 'cct16.txt').write_bytes(text.encode('utf-16'))
        (self.docs / 'cct32.txt').write_bytes(text.encode('utf-32'))
        (self.docs / 'sem-bom.txt').write_bytes(text.encode('utf-16-le'))
        utf16, utf32, no_bom = process([self.docs / n for n in ('cct16.txt', 'cct32.txt', 'sem-bom.txt')], self.out)
        self.assertIn(f'# Codificação: utf-16\n\n=== texto ===\n{text}\n', self.output(utf16))
        self.assertIn(f'# Codificação: utf-32\n\n=== texto ===\n{text}\n', self.output(utf32))
        self.assertEqual(utf16['caracteres'], len(text))
        self.assertIn('UTF-16 sem BOM', no_bom['erro'])

    def test_distinct_files_whose_names_differ_only_in_nfc_nfd_are_not_merged(self):
        nfc_name, nfd_name = (unicodedata.normalize(form, 'versão.txt') for form in ('NFC', 'NFD'))
        (self.docs / 'so-nfd').mkdir()
        alone = self.docs / 'so-nfd' / nfd_name
        alone.write_text('Versão única fictícia', encoding='utf-8')
        [entry] = process([alone, self.docs / 'so-nfd' / nfc_name], self.out)  # o mesmo arquivo, digitado em NFC
        self.assertIsNone(entry['erro'])
        (self.docs / nfc_name).write_text('Versão A fictícia', encoding='utf-8')
        (self.docs / nfd_name).write_text('Versão B fictícia', encoding='utf-8')
        if len([p for p in self.docs.iterdir() if p.is_file()]) != 2:
            self.skipTest('o sistema de arquivos não distingue NFC de NFD')
        code, out, err = run(self.docs / nfc_name, self.docs / nfd_name, '--saida', self.out)
        self.assertEqual(code, 1)
        self.assertEqual(json.loads(out)['com_erro'], 1)
        self.assertIn('mesmo nome em NFC', err)
        kept = next(e for e in self.inventory() if e['arquivo'] == unicodedata.normalize('NFC', str(self.docs / nfc_name)))
        self.assertIn('Versão A fictícia', self.output(kept))


if __name__ == '__main__':
    unittest.main()
